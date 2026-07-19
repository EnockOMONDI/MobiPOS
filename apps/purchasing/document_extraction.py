import csv
from io import BytesIO, StringIO
from pathlib import Path

from django.core.exceptions import ValidationError
from openpyxl import load_workbook

from .models import PurchaseDocumentExtractionStatus


TEXT_EXTENSIONS = {".csv", ".txt", ".tsv"}
SPREADSHEET_EXTENSIONS = {".xlsx", ".xlsm"}
MANUAL_REVIEW_EXTENSIONS = {".pdf", ".png", ".jpg", ".jpeg", ".webp", ".gif"}


def extract_purchase_document_text(upload):
    if not upload:
        raise ValidationError("Attach a supplier document before running extraction.")

    name = upload.name or ""
    extension = Path(name).suffix.lower()
    upload.open("rb")
    try:
        content = upload.read()
    finally:
        upload.close()

    if extension in TEXT_EXTENSIONS:
        text = content.decode("utf-8-sig", errors="replace")
        if extension in {".csv", ".tsv"}:
            delimiter = "\t" if extension == ".tsv" else ","
            reader = csv.reader(StringIO(text), delimiter=delimiter)
            rows = [" | ".join(cell.strip() for cell in row if cell is not None) for row in reader]
            text = "\n".join(row for row in rows if row.strip())
        return PurchaseDocumentExtractionStatus.EXTRACTED, text[:12000]

    if extension in SPREADSHEET_EXTENSIONS:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        lines = []
        for sheet in workbook.worksheets[:3]:
            lines.append(f"[{sheet.title}]")
            for row in sheet.iter_rows(max_row=80, max_col=12, values_only=True):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    lines.append(" | ".join(values))
        text = "\n".join(lines).strip()
        if text:
            return PurchaseDocumentExtractionStatus.EXTRACTED, text[:12000]
        return PurchaseDocumentExtractionStatus.MANUAL_REVIEW, "No readable spreadsheet cells were found."

    if extension in MANUAL_REVIEW_EXTENSIONS:
        return (
            PurchaseDocumentExtractionStatus.MANUAL_REVIEW,
            "This document appears to be scanned or image-based. Manual review is required unless a production OCR provider is configured.",
        )

    return (
        PurchaseDocumentExtractionStatus.FAILED,
        f"Unsupported supplier document format: {extension or 'unknown'}",
    )
