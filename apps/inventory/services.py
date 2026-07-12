from decimal import Decimal
import csv
from io import BytesIO, StringIO

from django.core.exceptions import ValidationError
from django.db import models, transaction
from django.db.models import F

from .models import SerialStatus, StockAdjustmentStatus, StockBalance, StockMovement, StockMovementType, StockUnit


@transaction.atomic
def post_stock_movement(*, organization, product, location, quantity, movement_type, actor=None, stock_unit=None, unit_cost=0, reference_type="", reference_id="", reason="", reversed_movement=None):
    quantity = Decimal(quantity)
    if quantity == 0:
        raise ValidationError("Stock movement quantity cannot be zero.")
    if product.organization_id != organization.id or location.organization_id != organization.id:
        raise ValidationError("Stock movement entities must belong to the same organization.")
    if stock_unit:
        if not product.is_serialized or abs(quantity) != 1:
            raise ValidationError("Serialized stock movements must have quantity 1 or -1.")
        if stock_unit.organization_id != organization.id or stock_unit.product_id != product.id:
            raise ValidationError("Stock unit must belong to the same organization and product.")
        if quantity < 0 and stock_unit.location_id != location.id:
            raise ValidationError("Stock unit is not at the movement source location.")
    elif product.is_serialized:
        raise ValidationError("Serialized stock movements require a stock unit.")
    balance, _ = StockBalance.objects.select_for_update().get_or_create(
        organization=organization, product=product, location=location
    )
    if balance.quantity + quantity < 0:
        raise ValidationError("Insufficient available stock.")
    StockBalance.objects.filter(pk=balance.pk).update(quantity=F("quantity") + quantity)
    return StockMovement.objects.create(
        organization=organization, movement_type=movement_type, product=product,
        stock_unit=stock_unit, location=location, quantity=quantity, unit_cost=unit_cost,
        reference_type=reference_type, reference_id=str(reference_id), reason=reason, actor=actor,
        reversed_movement=reversed_movement,
    )


@transaction.atomic
def complete_stock_adjustment(*, adjustment, actor):
    if adjustment.status != StockAdjustmentStatus.REQUESTED:
        raise ValidationError("Only requested stock adjustments can be completed.")
    movement = post_stock_movement(
        organization=adjustment.organization,
        product=adjustment.product,
        location=adjustment.location,
        quantity=adjustment.quantity,
        movement_type=StockMovementType.ADJUSTMENT,
        actor=actor,
        unit_cost=adjustment.product.cost_price,
        reference_type="stock_adjustment",
        reference_id=adjustment.id,
        reason=adjustment.reason,
    )
    adjustment.status = StockAdjustmentStatus.COMPLETED
    adjustment.approved_by = actor
    adjustment.movement = movement
    adjustment.save(update_fields=["status", "approved_by", "movement", "updated_at"])
    return adjustment


@transaction.atomic
def reverse_stock_movement(*, movement, actor, reason):
    movement = StockMovement.objects.select_for_update().get(pk=movement.pk)
    if movement.movement_type == StockMovementType.REVERSAL:
        raise ValidationError("A reversal movement cannot be reversed.")
    if movement.reversals.exists():
        raise ValidationError("This stock movement has already been reversed.")
    domain_owned_references = {
        "purchase_order",
        "supplier_return",
        "stock_transfer",
        "transfer_discrepancy",
        "sale",
        "sale_return",
        "sale_return_disposition",
        "repair_ticket",
    }
    if movement.reference_type in domain_owned_references:
        raise ValidationError("This movement belongs to a business transaction and must be corrected through its source workflow.")
    reversal = post_stock_movement(
        organization=movement.organization,
        product=movement.product,
        location=movement.location,
        quantity=-movement.quantity,
        movement_type=StockMovementType.REVERSAL,
        actor=actor,
        stock_unit=movement.stock_unit,
        unit_cost=movement.unit_cost,
        reference_type="stock_movement_reversal",
        reference_id=movement.id,
        reason=reason,
        reversed_movement=movement,
    )
    if movement.stock_unit:
        if reversal.quantity > 0:
            movement.stock_unit.status = SerialStatus.AVAILABLE
            movement.stock_unit.location = movement.location
        else:
            movement.stock_unit.status = SerialStatus.WRITTEN_OFF
            movement.stock_unit.location = None
        movement.stock_unit.save(update_fields=["status", "location", "updated_at"])
    return reversal


def _append_tabular_serial_rows(rows, table_rows, *, source):
    table_rows = list(table_rows)
    if not table_rows:
        return
    header = [str(value or "").strip() for value in table_rows[0]]
    normalized = {value.lower(): index for index, value in enumerate(header)}
    serial_index = normalized.get("serial_number", normalized.get("imei", normalized.get("serial")))
    secondary_index = normalized.get("secondary_serial", normalized.get("imei2", normalized.get("secondary_imei")))
    data_rows = table_rows[1:] if serial_index is not None else table_rows
    if serial_index is None:
        serial_index = 0
        secondary_index = 1
    for offset, row in enumerate(data_rows, start=2 if data_rows is not table_rows else 1):
        serial = str(row[serial_index] or "").strip() if len(row) > serial_index else ""
        secondary = str(row[secondary_index] or "").strip() if secondary_index is not None and len(row) > secondary_index else ""
        if serial:
            rows.append({"line": f"{source}:{offset}", "serial_number": serial, "secondary_serial": secondary})


def parse_serial_intake_rows(*, serial_numbers="", csv_file=None):
    rows = []
    for line_number, raw_line in enumerate((serial_numbers or "").replace(",", "\n").splitlines(), start=1):
        serial = raw_line.strip()
        if serial:
            rows.append({"line": f"text:{line_number}", "serial_number": serial, "secondary_serial": ""})
    if csv_file:
        filename = csv_file.name.lower()
        if filename.endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook

            workbook = load_workbook(BytesIO(csv_file.read()), read_only=True, data_only=True)
            worksheet = workbook.active
            _append_tabular_serial_rows(rows, worksheet.iter_rows(values_only=True), source="excel")
        else:
            content = csv_file.read().decode("utf-8-sig")
            reader = csv.reader(StringIO(content))
            _append_tabular_serial_rows(rows, reader, source="csv")
    return rows


@transaction.atomic
def batch_receive_serialized_stock(*, organization, product, location, rows, actor, unit_cost=0, reason=""):
    if product.organization_id != organization.id or location.organization_id != organization.id:
        raise ValidationError("Batch intake entities must belong to the same organization.")
    if not product.is_serialized:
        raise ValidationError("Batch intake requires a serialized product.")
    if not rows:
        raise ValidationError("No serial numbers were found.")
    cleaned_rows = []
    failures = []
    seen = {}
    for row in rows:
        serial = (row.get("serial_number") or "").strip()
        secondary = (row.get("secondary_serial") or "").strip()
        line = row.get("line", "")
        if not serial:
            failures.append({"line": line, "serial_number": serial, "error": "Serial number is required."})
            continue
        identifiers = [serial]
        if secondary:
            identifiers.append(secondary)
        if len(set(identifiers)) != len(identifiers):
            failures.append({"line": line, "serial_number": serial, "error": "Primary and secondary serial numbers must be different."})
            continue
        duplicate = next((identifier for identifier in identifiers if identifier in seen), None)
        if duplicate:
            failures.append({"line": line, "serial_number": serial, "error": f"Duplicate identifier in upload: {duplicate}."})
            continue
        for identifier in identifiers:
            seen[identifier] = line
        cleaned_rows.append({"line": line, "serial_number": serial, "secondary_serial": secondary})
    identifiers = [
        identifier
        for row in cleaned_rows
        for identifier in (row["serial_number"], row["secondary_serial"])
        if identifier
    ]
    existing = set(
        StockUnit.objects.filter(organization=organization).filter(
            models.Q(serial_number__in=identifiers) | models.Q(secondary_serial__in=identifiers)
        ).values_list("serial_number", flat=True)
    )
    existing_secondary = set(
        StockUnit.objects.filter(organization=organization, secondary_serial__in=identifiers)
        .exclude(secondary_serial="")
        .values_list("secondary_serial", flat=True)
    )
    existing |= existing_secondary
    accepted_rows = []
    for row in cleaned_rows:
        conflict = next((identifier for identifier in (row["serial_number"], row["secondary_serial"]) if identifier and identifier in existing), None)
        if conflict:
            failures.append({"line": row["line"], "serial_number": row["serial_number"], "error": f"Identifier already exists: {conflict}."})
            continue
        accepted_rows.append(row)
    created_units = []
    for row in accepted_rows:
        unit = StockUnit.objects.create(
            organization=organization,
            product=product,
            serial_number=row["serial_number"],
            secondary_serial=row["secondary_serial"],
            location=location,
            status=SerialStatus.AVAILABLE,
            unit_cost=unit_cost or 0,
        )
        post_stock_movement(
            organization=organization,
            product=product,
            location=location,
            quantity=1,
            movement_type=StockMovementType.OPENING,
            actor=actor,
            stock_unit=unit,
            unit_cost=unit_cost or 0,
            reference_type="batch_serial_intake",
            reference_id=unit.id,
            reason=reason or "Batch serialized stock intake",
        )
        created_units.append(unit)
    if not created_units and failures:
        raise ValidationError("No valid serial numbers were found.")
    return {"created_units": created_units, "failures": failures}
